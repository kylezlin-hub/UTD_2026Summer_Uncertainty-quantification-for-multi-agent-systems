"""Check Llama 8B debate accuracy — identify failures for Phase 2."""
import openpyxl
from pathlib import Path

RUNS_DIR = Path(r"C:\Proj1\Knowledge_vs_Reasoning\local\data\debate_llama_8b\runs")

workbooks = [
    ("mmlu", RUNS_DIR / "mmlu_seed_7" / "debate_local_7b_mmlu_s7.xlsx"),
    ("mmlu-pro", RUNS_DIR / "mmlu-pro_seed_7" / "debate_local_7b_mmlu-pro_s7.xlsx"),
    ("gpqa", RUNS_DIR / "gpqa_seed_7" / "debate_local_7b_gpqa_s7.xlsx"),
]

total_correct = 0
total_wrong = 0
failures_by_ds = {}

for ds, path in workbooks:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Round_State"]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]

    qno_idx = headers.index("question_no")
    round_idx = headers.index("round")
    agent_idx = headers.index("agent")
    answer_idx = headers.index("answer")
    correct_idx = headers.index("correct_answer")

    # Get final-round answers per question (majority vote)
    final_answers = {}
    correct_answers = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        qno = row[qno_idx]
        rnd = row[round_idx]
        ans = row[answer_idx]
        corr = row[correct_idx]
        correct_answers[qno] = corr
        if qno not in final_answers:
            final_answers[qno] = {}
        if qno not in final_answers or rnd > final_answers[qno].get("max_round", 0):
            final_answers[qno]["max_round"] = rnd
            final_answers[qno]["answers"] = []
        if rnd == final_answers[qno]["max_round"]:
            final_answers[qno]["answers"].append(ans)
        elif rnd > final_answers[qno]["max_round"]:
            final_answers[qno]["max_round"] = rnd
            final_answers[qno]["answers"] = [ans]

    wb.close()

    correct = 0
    wrong = 0
    failures = []
    for qno, data in final_answers.items():
        answers = data["answers"]
        from collections import Counter
        majority = Counter(answers).most_common(1)[0][0]
        if majority == correct_answers[qno]:
            correct += 1
        else:
            wrong += 1
            failures.append(qno)

    total_correct += correct
    total_wrong += wrong
    failures_by_ds[ds] = failures
    print(f"{ds}: {correct}/{correct+wrong} correct, {wrong} failures")

print(f"\nTotal: {total_correct}/{total_correct+total_wrong} correct, {total_wrong} failures")
print(f"\nFailures by dataset:")
for ds, f in failures_by_ds.items():
    print(f"  {ds}: {len(f)} failures")
print(f"\nTotal failures for Phase 2: {total_wrong}")
